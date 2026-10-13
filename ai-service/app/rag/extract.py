import io
import os
import re
import shutil
import zipfile

ALLOWED_EXTENSIONS = {".txt", ".md", ".pdf", ".docx", ".xlsx", ".pptx", ".png", ".jpg", ".jpeg"}
IMAGE_EXTENSIONS = {".png", ".jpg", ".jpeg"}
MAX_PDF_PAGES = 300
MAX_DOCX_UNZIPPED = 50 * 1024 * 1024  # chống zip bomb (docx/xlsx/pptx đều là zip)
MAX_SHEET_ROWS = 20000                # tổng số dòng đọc từ mọi sheet của một file Excel
MAX_OCR_PAGES = 30                    # PDF scan: chỉ OCR tối đa ngần này trang đầu
OCR_TIMEOUT_SECONDS = 60              # mỗi ảnh/trang


class ExtractError(ValueError):
    """File hợp lệ về đuôi nhưng không đọc được nội dung."""


def extension(filename: str) -> str:
    name = (filename or "").lower()
    dot = name.rfind(".")
    return name[dot:] if dot >= 0 else ""


def _clean(text: str) -> str:
    # PostgreSQL TEXT không chấp nhận ký tự NUL.
    text = text.replace("\x00", "")
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    text = re.sub(r"[ \t]+\n", "\n", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def _decode_text(data: bytes) -> str:
    # utf-16 chỉ khi có BOM: không có BOM thì byte bất kỳ đều "giải mã được" thành rác.
    if data.startswith((b"\xff\xfe", b"\xfe\xff")):
        return data.decode("utf-16")
    for encoding in ("utf-8-sig", "cp1258"):
        try:
            return data.decode(encoding)
        except UnicodeError:
            continue
    return data.decode("utf-8", errors="replace")


def _read_pdf(data: bytes) -> str:
    from pypdf import PdfReader

    try:
        reader = PdfReader(io.BytesIO(data))
        if reader.is_encrypted and not reader.decrypt(""):
            raise ExtractError("PDF được bảo vệ bằng mật khẩu.")
        if len(reader.pages) > MAX_PDF_PAGES:
            raise ExtractError(f"PDF quá {MAX_PDF_PAGES} trang.")
        pages = [(page.extract_text() or "") for page in reader.pages]
    except ExtractError:
        raise
    except Exception as exc:
        raise ExtractError("Không đọc được file PDF (hỏng hoặc sai định dạng).") from exc

    return "\n\n".join(pages)


def _read_docx(data: bytes) -> str:
    from docx import Document

    try:
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            if sum(i.file_size for i in archive.infolist()) > MAX_DOCX_UNZIPPED:
                raise ExtractError("File DOCX giải nén quá lớn.")
        document = Document(io.BytesIO(data))
    except ExtractError:
        raise
    except Exception as exc:
        raise ExtractError("Không đọc được file DOCX (hỏng hoặc sai định dạng).") from exc

    blocks = [p.text for p in document.paragraphs if p.text.strip()]
    for table in document.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                blocks.append(" | ".join(cells))
    return "\n\n".join(blocks)


def _check_zip(data: bytes, label: str) -> None:
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            if sum(i.file_size for i in archive.infolist()) > MAX_DOCX_UNZIPPED:
                raise ExtractError(f"File {label} giải nén quá lớn.")
    except ExtractError:
        raise
    except Exception as exc:
        raise ExtractError(f"Không đọc được file {label} (hỏng hoặc sai định dạng).") from exc


def _cell(value) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return " ".join(str(value).split())


def _read_xlsx(data: bytes) -> str:
    """Mỗi sheet thành một khối; dòng đầu là tiêu đề cột thì mỗi dòng dữ liệu viết dạng 'Cột: giá trị'
    để câu hỏi như 'doanh thu tháng 3' khớp được cả tên cột lẫn giá trị."""
    from openpyxl import load_workbook

    _check_zip(data, "XLSX")
    try:
        workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception as exc:
        raise ExtractError("Không đọc được file XLSX (hỏng hoặc sai định dạng).") from exc

    blocks, total = [], 0
    try:
        for sheet in workbook.worksheets:
            lines, header = [], None
            for row in sheet.iter_rows(values_only=True):
                cells = [_cell(v) for v in row]
                if not any(cells):
                    continue
                total += 1
                if total > MAX_SHEET_ROWS:
                    raise ExtractError(f"File Excel quá {MAX_SHEET_ROWS} dòng dữ liệu. Hãy tách nhỏ file.")
                if header is None:
                    named = [c for c in cells if c]
                    header = cells if len(named) >= 2 and all(not c.replace('.', '', 1).isdigit() for c in named) else []
                    if header:
                        lines.append(" | ".join(named))
                        continue
                if header:
                    pairs = [f"{h}: {c}" if h else c for h, c in zip(header, cells) if c]
                    lines.append(" | ".join(pairs))
                else:
                    lines.append(" | ".join(c for c in cells if c))
            if lines:
                blocks.append(f"## Sheet {sheet.title}\n" + "\n".join(lines))
    finally:
        workbook.close()
    return "\n\n".join(blocks)


def _read_pptx(data: bytes) -> str:
    from pptx import Presentation

    _check_zip(data, "PPTX")
    try:
        deck = Presentation(io.BytesIO(data))
    except Exception as exc:
        raise ExtractError("Không đọc được file PPTX (hỏng hoặc sai định dạng).") from exc

    blocks = []
    for number, slide in enumerate(deck.slides, start=1):
        title = slide.shapes.title.text.strip() if slide.shapes.title is not None and slide.shapes.title.has_text_frame else ""
        lines = [f"## Slide {number}" + (f": {title}" if title else "")]
        for shape in slide.shapes:
            if shape.has_text_frame and shape is not slide.shapes.title:
                lines += [p.text.strip() for p in shape.text_frame.paragraphs if p.text.strip()]
            if getattr(shape, "has_table", False) and shape.has_table:
                for row in shape.table.rows:
                    cells = [c.text.strip() for c in row.cells if c.text.strip()]
                    if cells:
                        lines.append(" | ".join(cells))
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame is not None:
            notes = slide.notes_slide.notes_text_frame.text.strip()
            if notes:
                lines.append("Ghi chú: " + notes)
        if len(lines) > 1:
            blocks.append("\n".join(lines))
    return "\n\n".join(blocks)


def ocr_available() -> bool:
    return os.getenv("OCR_ENABLED", "true").strip().lower() not in {"0", "false", "no"} and shutil.which("tesseract") is not None


def _ocr_langs() -> str:
    """Ngôn ngữ OCR: OCR_LANGS (mặc định vie+eng) lọc theo gói ngôn ngữ thực sự đã cài trong tesseract."""
    import pytesseract

    wanted = [l for l in os.getenv("OCR_LANGS", "vie+eng").replace(",", "+").split("+") if l]
    installed = set(pytesseract.get_languages(config=""))
    usable = [l for l in wanted if l in installed]
    if not usable:
        raise ExtractError("OCR chưa có gói ngôn ngữ phù hợp (cần tesseract-ocr-vie hoặc tesseract-ocr-eng).")
    return "+".join(usable)


def _ocr_image(image) -> str:
    import pytesseract

    if image.mode not in ("RGB", "L"):
        image = image.convert("RGB")
    try:
        return pytesseract.image_to_string(image, lang=_ocr_langs(), timeout=OCR_TIMEOUT_SECONDS)
    except RuntimeError as exc:  # hết thời gian
        raise ExtractError("OCR quá thời gian xử lý một ảnh/trang.") from exc


def _need_ocr() -> None:
    if not ocr_available():
        raise ExtractError("OCR chưa được bật trên máy chủ (cần cài tesseract-ocr; OCR_ENABLED=true).")


def _read_image(data: bytes) -> str:
    from PIL import Image

    _need_ocr()
    try:
        with Image.open(io.BytesIO(data)) as image:
            image.load()
            return _ocr_image(image)
    except ExtractError:
        raise
    except Exception as exc:
        raise ExtractError("Không đọc được ảnh (hỏng, sai định dạng hoặc quá lớn).") from exc


def _ocr_pdf(data: bytes) -> str:
    """PDF scan (không có lớp chữ): vẽ từng trang thành ảnh rồi OCR."""
    import pypdfium2

    _need_ocr()
    try:
        pdf = pypdfium2.PdfDocument(data)
        texts = []
        for index in range(min(len(pdf), MAX_OCR_PAGES)):
            texts.append(_ocr_image(pdf[index].render(scale=2).to_pil()))
        return "\n\n".join(texts)
    except ExtractError:
        raise
    except Exception as exc:
        raise ExtractError("Không OCR được file PDF scan.") from exc


def extract_text(filename: str, data: bytes) -> str:
    ext = extension(filename)

    if ext in (".txt", ".md"):
        text = _decode_text(data)
    elif ext == ".pdf":
        if not data.startswith(b"%PDF"):
            raise ExtractError("Nội dung không phải PDF.")
        text = _read_pdf(data)
        if not _clean(text):
            text = _ocr_pdf(data)  # PDF scan: không có lớp chữ nên đọc bằng OCR
    elif ext == ".docx":
        if not data.startswith(b"PK"):
            raise ExtractError("Nội dung không phải DOCX.")
        text = _read_docx(data)
    elif ext == ".xlsx":
        if not data.startswith(b"PK"):
            raise ExtractError("Nội dung không phải XLSX.")
        text = _read_xlsx(data)
    elif ext == ".pptx":
        if not data.startswith(b"PK"):
            raise ExtractError("Nội dung không phải PPTX.")
        text = _read_pptx(data)
    elif ext in IMAGE_EXTENSIONS:
        magic = (b"\x89PNG",) if ext == ".png" else (b"\xff\xd8",)
        if not data.startswith(magic):
            raise ExtractError("Nội dung không đúng định dạng ảnh.")
        text = _read_image(data)
    else:
        raise ExtractError(f"Định dạng {ext or '(không đuôi)'} chưa được hỗ trợ.")

    text = _clean(text)
    if not text:
        raise ExtractError(
            "Không trích xuất được văn bản (file không có chữ, hoặc OCR không đọc được ảnh)."
        )
    return text
