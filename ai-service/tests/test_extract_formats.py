"""Trích văn bản từ Excel, PowerPoint, ảnh (OCR) và PDF scan."""
import io

import pytest

from app.rag import extract
from app.rag.extract import ExtractError, extract_text

needs_ocr = pytest.mark.skipif(not extract.ocr_available(), reason="chưa cài tesseract")


def make_xlsx():
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    ws.title = "Doanh thu"
    ws.append(["Tháng", "Doanh thu", "Khu vực"])
    ws.append(["Tháng 3", 1250000, "Miền Nam"])
    ws.append([None, None, None])
    ws.append(["Tháng 4", 1980000.0, "Miền Bắc"])
    ws2 = wb.create_sheet("Ghi chú")
    ws2.append(["Số liệu chưa kiểm toán"])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def make_pptx():
    from pptx import Presentation
    from pptx.util import Inches
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = "Kế hoạch bảo trì"
    slide.placeholders[1].text = "Bảo dưỡng máy nén mỗi quý\nThay lọc dầu hằng năm"
    slide.notes_slide.notes_text_frame.text = "Nhớ báo trước cho ca đêm"
    blank = prs.slides.add_slide(prs.slide_layouts[6])
    rows = blank.shapes.add_table(2, 2, Inches(1), Inches(1), Inches(4), Inches(1)).table
    rows.cell(0, 0).text, rows.cell(0, 1).text = "Hạng mục", "Chi phí"
    rows.cell(1, 0).text, rows.cell(1, 1).text = "Lọc dầu", "500k"
    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


def make_image(text="Invoice total 2026 refund policy", size=64, fmt="PNG"):
    from PIL import Image, ImageDraw, ImageFont
    img = Image.new("RGB", (1600, 220), "white")
    ImageDraw.Draw(img).text((30, 60), text, fill="black", font=ImageFont.load_default(size=size))
    buf = io.BytesIO()
    img.save(buf, fmt)
    return buf.getvalue(), img


def test_xlsx_uses_headers_and_skips_blank_rows():
    text = extract_text("bang.xlsx", make_xlsx())
    assert "## Sheet Doanh thu" in text and "## Sheet Ghi chú" in text
    assert "Tháng: Tháng 3 | Doanh thu: 1250000 | Khu vực: Miền Nam" in text
    assert "Doanh thu: 1980000" in text  # số thực nguyên bỏ phần .0
    assert "Số liệu chưa kiểm toán" in text


def test_pptx_slides_title_body_table_and_notes():
    text = extract_text("slide.pptx", make_pptx())
    assert "## Slide 1: Kế hoạch bảo trì" in text
    assert "Bảo dưỡng máy nén mỗi quý" in text and "Thay lọc dầu hằng năm" in text
    assert "Ghi chú: Nhớ báo trước cho ca đêm" in text
    assert "Hạng mục | Chi phí" in text and "Lọc dầu | 500k" in text


@pytest.mark.parametrize("name,data", [
    ("hong.xlsx", b"PK\x03\x04 khong phai xlsx that"),
    ("hong.pptx", b"PK\x03\x04 khong phai pptx that"),
    ("gia.xlsx", b"not a zip"),
    ("gia.png", b"not a png"),
    ("gia.jpg", b"\x89PNG sai duoi"),
])
def test_corrupt_or_fake_files_raise_extract_error(name, data):
    with pytest.raises(ExtractError):
        extract_text(name, data)


def test_empty_workbook_has_clear_message():
    from openpyxl import Workbook
    buf = io.BytesIO()
    Workbook().save(buf)
    with pytest.raises(ExtractError, match="Không trích xuất"):
        extract_text("trong.xlsx", buf.getvalue())


def test_image_without_ocr_gives_clear_error(monkeypatch):
    monkeypatch.setenv("OCR_ENABLED", "false")
    data, _ = make_image()
    with pytest.raises(ExtractError, match="OCR chưa được bật"):
        extract_text("anh.png", data)


@needs_ocr
@pytest.mark.parametrize("fmt,name", [("PNG", "anh.png"), ("JPEG", "anh.jpg")])
def test_image_ocr_reads_text(fmt, name):
    data, _ = make_image(fmt=fmt)
    text = extract_text(name, data).lower()
    assert "invoice" in text and "refund" in text


@needs_ocr
def test_scanned_pdf_falls_back_to_ocr():
    _, img = make_image("Scanned contract number 7788")
    buf = io.BytesIO()
    img.save(buf, "PDF")  # PDF chỉ chứa ảnh, không có lớp chữ
    text = extract_text("scan.pdf", buf.getvalue()).lower()
    assert "contract" in text and "7788" in text
