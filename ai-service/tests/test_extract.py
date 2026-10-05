import io
import zipfile

import pytest

from app.extract import ExtractError, extension, extract_text


def test_extension():
    assert extension("A.PDF") == ".pdf"
    assert extension("noext") == ""
    assert extension("a.tar.gz") == ".gz"


def test_text_utf8_and_cp1258():
    assert extract_text("a.txt", "Xin chào\r\n\r\n\r\nbạn".encode("utf-8")) == "Xin chào\n\nbạn"
    # File Windows tiếng Việt cũ (không phải UTF-8) vẫn đọc được.
    assert extract_text("a.txt", "Xin chào".encode("cp1258")) == "Xin chào"


def test_utf16_with_bom():
    assert extract_text("a.md", "Xin chào".encode("utf-16")) == "Xin chào"


def test_nul_bytes_are_removed():
    assert "\x00" not in extract_text("a.txt", b"abc\x00def")


def test_empty_file_is_rejected():
    with pytest.raises(ExtractError):
        extract_text("a.txt", b"   \n\n ")


def test_unsupported_extension():
    with pytest.raises(ExtractError):
        extract_text("a.exe", b"MZ")


def test_magic_bytes_are_checked():
    with pytest.raises(ExtractError):
        extract_text("fake.pdf", b"not a pdf")
    with pytest.raises(ExtractError):
        extract_text("fake.docx", b"not a docx")


def test_corrupt_pdf_and_docx():
    with pytest.raises(ExtractError):
        extract_text("bad.pdf", b"%PDF-1.4 garbage")
    with pytest.raises(ExtractError):
        extract_text("bad.docx", b"PK\x03\x04garbage")


def test_docx_roundtrip():
    from docx import Document

    doc = Document()
    doc.add_paragraph("Đoạn văn thử nghiệm")
    table = doc.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "A"
    table.rows[0].cells[1].text = "B"
    buffer = io.BytesIO()
    doc.save(buffer)

    text = extract_text("x.docx", buffer.getvalue())
    assert "Đoạn văn thử nghiệm" in text
    assert "A | B" in text


def test_docx_zip_bomb_is_rejected(monkeypatch):
    from app import extract

    monkeypatch.setattr(extract, "MAX_DOCX_UNZIPPED", 100)
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED) as z:
        z.writestr("word/document.xml", "a" * 10_000)
    with pytest.raises(ExtractError):
        extract_text("bomb.docx", buffer.getvalue())